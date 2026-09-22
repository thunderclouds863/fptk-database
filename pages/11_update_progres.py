# pages/11_update_progres.py
import streamlit as st
import pandas as pd
from datetime import datetime, date, timedelta
from core.database import get_db
from core.models import FPTK, User, RecruitmentProgress
from core.auth import get_current_user, is_admin
from core.template_manager import (
    save_template, get_active_template, get_template_bytes, get_template_history
)
import time

from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def get_current_week():
    today = date.today()
    iso = today.isocalendar()
    return iso[1], iso[0]


def get_week_label(week_num, year):
    return f"Week {week_num}, {year}"


def get_week_range(week_num, year):
    try:
        jan4 = date(year, 1, 4)
        start_of_year_week = jan4 - timedelta(days=jan4.isoweekday() - 1)
        week_start = start_of_year_week + timedelta(weeks=week_num - 1)
        week_end = week_start + timedelta(days=6)
        return week_start, week_end
    except Exception:
        return None, None


@st.cache_data(ttl=300, show_spinner=False)
def _load_filter_options():
    db_local = next(get_db())
    try:
        pic_opts = sorted([
            r[0] for r in db_local.query(FPTK.pic_recruiter)
            .filter(FPTK.pic_recruiter.isnot(None), FPTK.pic_recruiter != "")
            .distinct().all() if r[0]
        ])
        bu_opts = sorted([
            r[0] for r in db_local.query(FPTK.business_unit)
            .filter(FPTK.business_unit.isnot(None), FPTK.business_unit != "")
            .distinct().all() if r[0]
        ])
        dir_opts = sorted([
            r[0] for r in db_local.query(FPTK.direktorat)
            .filter(FPTK.direktorat.isnot(None), FPTK.direktorat != "")
            .distinct().all() if r[0]
        ])
        kat_opts = sorted([
            r[0] for r in db_local.query(FPTK.filter_kategorisasi_fptk)
            .filter(FPTK.filter_kategorisasi_fptk.isnot(None),
                    FPTK.filter_kategorisasi_fptk != "")
            .distinct().all() if r[0]
        ])
        divisi_opts = sorted([
            r[0] for r in db_local.query(FPTK.divisi)
            .filter(FPTK.divisi.isnot(None), FPTK.divisi != "")
            .distinct().all() if r[0]
        ])
        dept_opts = sorted([
            r[0] for r in db_local.query(FPTK.department)
            .filter(FPTK.department.isnot(None), FPTK.department != "")
            .distinct().all() if r[0]
        ])
        return {
            "pic_options": pic_opts,
            "bu_options": bu_opts,
            "direktorat_options": dir_opts,
            "filter_kategorisasi_options": kat_opts,
            "divisi_options": divisi_opts,
            "dept_options": dept_opts,
        }
    except Exception as e:
        st.error(f"Gagal load filter options: {e}")
        return {
            "pic_options": [], "bu_options": [], "direktorat_options": [],
            "filter_kategorisasi_options": [], "divisi_options": [], "dept_options": [],
        }
    finally:
        db_local.close()


def generate_progress_recruitment_excel(df):
    output = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Update Progress Recruitment"

    headers = [
        'No', 'Tanggal FPTK', 'Posisi', 'Level',
        'Business Unit', 'Filter Kategorisasi', 'SLA Target Pemenuhan',
        'PIC TA', 'Vacancy', 'Status', 'Recruitment Update'
    ]

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

    for row_idx, row in df.iterrows():
        for col_idx, value in enumerate(row, start=1):
            if pd.isna(value):
                value = ""
            cell = ws.cell(row=row_idx + 2, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(
                left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin')
            )

    column_widths = {
        'A': 6, 'B': 14, 'C': 38, 'D': 7, 'E': 30,
        'F': 14, 'G': 14, 'H': 10, 'I': 12, 'J': 12, 'K': 75,
    }
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[1].height = 35
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{len(df) + 1}"

    wb.save(output)
    output.seek(0)
    return output


def upsert_progress(db, fptk, week_num, year, progress_text, next_action=""):
    existing = db.query(RecruitmentProgress).filter(
        RecruitmentProgress.fptk_id == fptk.id,
        RecruitmentProgress.week_number == week_num,
        RecruitmentProgress.year == year
    ).first()

    week_label = f"Week {week_num}, {year}"

    if existing:
        existing.progress_this_week = progress_text
        if next_action:
            existing.next_action = next_action
        existing.updated_at = datetime.now()
        return existing, "updated"
    else:
        new_progress = RecruitmentProgress(
            fptk_id=fptk.id,
            kode_unik=fptk.kode_unik,
            posisi=fptk.posisi,
            pic_recruiter=fptk.pic_recruiter,
            week_number=week_num,
            year=year,
            week_label=week_label,
            progress_this_week=progress_text,
            next_action=next_action,
            status="SUBMITTED",
            created_by_name="Excel Upload",
        )
        db.add(new_progress)
        return new_progress, "created"


# ============================================================
# TAB 1: UPDATE MANUAL
# ============================================================

def tab_update_manual(db, user, admin, current_week, current_year, filter_opts):
    query = db.query(FPTK).filter(FPTK.status == "OP")

    if st.session_state.get("search_progres"):
        s = st.session_state.search_progres.strip()
        query = query.filter(
            (FPTK.kode_unik.ilike(f"%{s}%")) | (FPTK.posisi.ilike(f"%{s}%"))
        )

    if st.session_state.get("pic_progres") and st.session_state.pic_progres != "Semua":
        query = query.filter(FPTK.pic_recruiter == st.session_state.pic_progres)
    if st.session_state.get("bu_progres") and st.session_state.bu_progres != "Semua":
        query = query.filter(FPTK.business_unit == st.session_state.bu_progres)
    if st.session_state.get("dir_progres") and st.session_state.dir_progres != "Semua":
        query = query.filter(FPTK.direktorat == st.session_state.dir_progres)
    if st.session_state.get("divisi_progres") and st.session_state.divisi_progres != "Semua":
        query = query.filter(FPTK.divisi == st.session_state.divisi_progres)
    if st.session_state.get("dept_progres") and st.session_state.dept_progres != "Semua":
        query = query.filter(FPTK.department == st.session_state.dept_progres)
    if st.session_state.get("kat_progres") and st.session_state.kat_progres != "Semua":
        query = query.filter(FPTK.filter_kategorisasi_fptk == st.session_state.kat_progres)
    if st.session_state.get("level_progres") and st.session_state.level_progres != "Semua":
        query = query.filter(FPTK.level_fptk == st.session_state.level_progres)
    if st.session_state.get("show_mine_progres") and not admin:
        query = query.filter(FPTK.pic_recruiter == user.pic_recruiter)

    query = query.order_by(FPTK.fptk_date_real.desc())
    total = query.count()

    st.markdown("### 📈 Statistik")

    col1, col2, col3, col4 = st.columns(4)

    existing_progress = db.query(RecruitmentProgress).filter(
        RecruitmentProgress.week_number == current_week,
        RecruitmentProgress.year == current_year
    ).all()
    progress_fptk_ids = set([p.fptk_id for p in existing_progress])

    already_updated = 0
    if progress_fptk_ids:
        already_updated = query.filter(FPTK.id.in_(progress_fptk_ids)).count()
    belum_update = total - already_updated

    col1.metric("Total FPTK OP", total)
    col2.metric("✅ Sudah Update", already_updated)
    col3.metric("⏳ Belum Update", belum_update)
    col4.metric("📊 Progress", f"{(already_updated/total*100):.0f}%" if total > 0 else "0%")

    st.markdown("---")

    if total == 0:
        st.info("Tidak ada FPTK dengan status **OP** yang sesuai filter.")
        return

    page_size = st.number_input("Baris per halaman", min_value=10, max_value=200, value=50, key="up_page_size")
    page = st.number_input("Halaman", min_value=1, max_value=max(1, (total + page_size - 1) // page_size), value=1, key="up_page")
    offset = (page - 1) * page_size

    df = pd.read_sql(query.limit(page_size).offset(offset).statement, db.bind)

    if not df.empty:
        fptk_ids_page = df['id'].tolist()
        progress_map = {p.fptk_id: p for p in existing_progress if p.fptk_id in fptk_ids_page}

        display_data = []
        for _, row in df.iterrows():
            has_progress = row['id'] in progress_map
            display_data.append({
                "Kode Unik": row.get('kode_unik', ''),
                "Posisi": (row.get('posisi', '') or '')[:60],
                "PIC": row.get('pic_recruiter', ''),
                "BU": (row.get('business_unit', '') or '')[:30],
                "Level": row.get('level_fptk', ''),
                "Status Update": "✅ Sudah" if has_progress else "⏳ Belum",
            })

        st.dataframe(pd.DataFrame(display_data), use_container_width=True, height=400, hide_index=True)

    st.markdown("---")
    st.markdown("### ✏️ Update Progress FPTK")

    df_all = pd.read_sql(query.statement, db.bind)
    if df_all.empty:
        return

    select_options = {}
    for _, row in df_all.iterrows():
        kode = row.get('kode_unik', '')
        posisi = row.get('posisi', '')
        display = f"{kode} | {posisi[:50]}" if len(str(posisi)) > 50 else f"{kode} | {posisi}"
        select_options[display] = row.get('id')

    selected_display = st.selectbox("Pilih FPTK", list(select_options.keys()), key="up_fptk_select")
    selected_id = select_options[selected_display]

    detail = db.query(FPTK).filter(FPTK.id == selected_id).first()
    if not detail:
        st.error("FPTK tidak ditemukan!")
        return

    can_edit = admin or (detail.pic_recruiter == user.pic_recruiter)
    if not can_edit:
        st.warning("Anda hanya bisa update progress FPTK milik PIC Anda sendiri.")
        return

    col1, col2, col3 = st.columns(3)
    with col1:
        st.markdown(f"**Kode Unik:** {detail.kode_unik}")
        st.markdown(f"**Posisi:** {detail.posisi}")
    with col2:
        st.markdown(f"**PIC:** {detail.pic_recruiter}")
        st.markdown(f"**BU:** {detail.business_unit}")
    with col3:
        st.markdown(f"**Level:** {detail.level_fptk}")
        st.markdown(f"**Status:** {detail.status}")

    existing = db.query(RecruitmentProgress).filter(
        RecruitmentProgress.fptk_id == selected_id,
        RecruitmentProgress.week_number == current_week,
        RecruitmentProgress.year == current_year
    ).first()

    default_progress = existing.progress_this_week if existing else ""
    default_next = existing.next_action if existing else ""

    form_key_progress = f"form_progress_text_{selected_id}"
    form_key_next = f"form_next_text_{selected_id}"

    if form_key_progress not in st.session_state:
        st.session_state[form_key_progress] = default_progress
    if form_key_next not in st.session_state:
        st.session_state[form_key_next] = default_next

    with st.form(f"form_progress_{selected_id}", clear_on_submit=False):
        st.markdown(f"#### 📅 Update untuk **{get_week_label(current_week, current_year)}**")

        if existing:
            st.info(f"✏️ Sudah ada update sebelumnya.")

        progress_this_week = st.text_area(
            "📝 **Progress Week Ini**",
            value=st.session_state[form_key_progress],
            key=f"text_area_progress_{selected_id}",
            placeholder="Contoh:\n> Send 25 CV\n> Shortlisted 9 kandidat\n> HR Interview 5 kandidat",
            height=200
        )

        next_action = st.text_area(
            "➡️ **Next Action Week Depan**",
            value=st.session_state[form_key_next],
            key=f"text_area_next_{selected_id}",
            placeholder="Contoh:\n> Send 10 kandidat baru\n> Follow up user interview",
            height=150
        )

        submit = st.form_submit_button("💾 Simpan Progress", type="primary", use_container_width=True)

        if submit:
            if not progress_this_week.strip():
                st.error("❌ Progress Week Ini wajib diisi!")
            elif not next_action.strip():
                st.error("❌ Next Action wajib diisi!")
            else:
                try:
                    progress, action = upsert_progress(
                        db, detail, current_week, current_year,
                        progress_this_week.strip(), next_action.strip()
                    )
                    progress.created_by = user.id
                    progress.created_by_name = user.display_name or user.username
                    db.commit()

                    st.session_state[form_key_progress] = ""
                    st.session_state[form_key_next] = ""

                    st.cache_data.clear()
                    st.success(f"✅ Progress berhasil disimpan!")
                    time.sleep(0.5)
                    st.rerun()
                except Exception as e:
                    st.error(f"❌ Error: {str(e)}")
                    db.rollback()

    st.markdown("---")
    st.markdown(f"### 📜 History Progress — {detail.kode_unik}")

    history = db.query(RecruitmentProgress).filter(
        RecruitmentProgress.fptk_id == selected_id
    ).order_by(
        RecruitmentProgress.year.desc(),
        RecruitmentProgress.week_number.desc()
    ).limit(20).all()

    if not history:
        st.info("Belum ada history progress.")
    else:
        for h in history:
            with st.expander(
                f"📅 **{h.week_label}** — {h.created_by_name} pada {h.created_at.strftime('%d/%m/%Y %H:%M') if h.created_at else '-'}",
                expanded=(h.week_number == current_week and h.year == current_year)
            ):
                col1, col2 = st.columns(2)
                with col1:
                    st.markdown("**📝 Progress:**")
                    st.info(h.progress_this_week or "-")
                with col2:
                    st.markdown("**➡️ Next Action:**")
                    st.success(h.next_action or "-")

                if admin:
                    if st.button(f"🗑️ Hapus", key=f"del_prog_{h.id}"):
                        try:
                            db.delete(h)
                            db.commit()
                            st.rerun()
                        except Exception as e:
                            st.error(f"❌ {str(e)}")
                            db.rollback()


# ============================================================
# TAB 2: UPLOAD EXCEL MASSAL
# ============================================================

def tab_upload_excel(db, user, admin, current_week, current_year):
    st.markdown("### 📤 Upload Excel Progress Massal")
    st.caption("Upload file Excel untuk import progress sekaligus banyak.")

    # ============================================================
    # ADMIN: TEMPLATE MANAGEMENT
    # ============================================================
    if admin:
        with st.expander("⚙️ Admin - Kelola Template Excel Progress", expanded=False):
            st.markdown("#### 📋 Template Aktif")

            active_template = get_active_template(db, template_type="PROGRESS")

            if active_template:
                template_bytes = get_template_bytes(active_template)
                col1, col2 = st.columns([2, 1])

                with col1:
                    st.markdown(f"**File:** `{active_template.file_name}`")
                    st.caption(f"Versi: **{active_template.version}** | Diupload: {active_template.created_at.strftime('%d/%m/%Y %H:%M') if active_template.created_at else '-'}")
                    st.caption(f"Ukuran: {len(template_bytes) / 1024:.1f} KB")

                with col2:
                    st.download_button(
                        "📥 Download Template",
                        template_bytes,
                        active_template.file_name,
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True,
                        key="admin_dl_progress_template"
                    )
            else:
                st.warning("⚠️ Belum ada template aktif untuk Update Progress Recruitment.")

            st.markdown("---")
            st.markdown("#### 📤 Upload Template Baru")
            st.caption("Upload file Excel template baru. Template lama otomatis jadi non-aktif. Versi auto increment.")

            new_template_file = st.file_uploader(
                "Pilih file Excel template",
                type=["xlsx", "xlsm"],
                key="admin_upload_progress_template"
            )

            if new_template_file:
                # Validasi: cek header
                try:
                    test_df = pd.read_excel(new_template_file, nrows=0)
                    cols_lower = [str(c).lower().strip() for c in test_df.columns]

                    has_kode = any("kode" in c and "unik" in c for c in cols_lower)
                    has_progress = any(("recruitment" in c and "update" in c) or "progress" in c for c in cols_lower)

                    if has_kode and has_progress:
                        st.success(f"✅ Template valid! Kolom: {', '.join(test_df.columns.tolist())}")

                        # Preview 5 baris pertama
                        preview_df = pd.read_excel(new_template_file, nrows=5)
                        st.markdown("**Preview 5 baris pertama:**")
                        st.dataframe(preview_df, use_container_width=True)

                        # Reset file pointer
                        new_template_file.seek(0)

                        col_btn1, col_btn2 = st.columns(2)
                        with col_btn1:
                            if st.button("💾 Simpan Template", type="primary", use_container_width=True, key="admin_save_progress_template"):
                                try:
                                    saved = save_template(db, new_template_file, user.id, template_type="PROGRESS")
                                    st.success(f"✅ Template berhasil disimpan! Versi: {saved.version}")
                                    st.cache_data.clear()
                                    time.sleep(0.5)
                                    st.rerun()
                                except Exception as e:
                                    st.error(f"❌ Error: {str(e)}")
                    else:
                        st.error("❌ Template tidak valid!")
                        missing = []
                        if not has_kode:
                            missing.append("**Kode Unik**")
                        if not has_progress:
                            missing.append("**Recruitment Update** atau **Progress**")
                        st.warning(f"Kolom yang wajib ada: {', '.join(missing)}")
                        st.caption(f"Kolom yang ditemukan: {', '.join(test_df.columns.tolist())}")

                except Exception as e:
                    st.error(f"❌ Gagal baca file: {str(e)}")

            st.markdown("---")
            st.markdown("#### 📜 History Template")

            history = get_template_history(db, template_type="PROGRESS", limit=10)

            if history:
                hist_data = []
                for h in history:
                    hist_data.append({
                        "Versi": h.version,
                        "File": h.file_name,
                        "Status": "✅ Aktif" if h.is_active else "⏸️ Non-aktif",
                        "Tanggal": h.created_at.strftime("%d/%m/%Y %H:%M") if h.created_at else "-",
                    })
                st.dataframe(pd.DataFrame(hist_data), use_container_width=True, hide_index=True)
            else:
                st.info("Belum ada history template.")

    # ============================================================
    # USER: DOWNLOAD TEMPLATE + PETUNJUK
    # ============================================================
    st.markdown("---")
    st.markdown("#### 📥 Download Template")
    st.caption("Download template untuk memudahkan pengisian progress recruitment.")

    active_template = get_active_template(db, template_type="PROGRESS")

    if active_template:
        template_bytes = get_template_bytes(active_template)
        col1, col2 = st.columns([3, 1])

        with col1:
            st.info(f"📋 Template aktif: **{active_template.file_name}** (v{active_template.version})")

        with col2:
            st.download_button(
                "📥 Download Template",
                template_bytes,
                active_template.file_name,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                key="user_dl_progress_template"
            )
    else:
        st.warning("⚠️ Belum ada template. Hubungi Admin.")

    # ============================================================
    # PETUNJUK
    # ============================================================
    with st.expander("📋 Petunjuk Penggunaan Template", expanded=False):
        st.markdown("""
        ### 🎯 Cara Menggunakan Template

        **Langkah 1: Download Template**
        - Klik tombol **📥 Download Template** di atas
        - Template akan berisi kolom:
          - **Kode Unik** — Kode unik FPTK (wajib)
          - **Recruitment Update** — Progress recruitment (wajib)
          - **Next Action** — Next action (opsional)

        **Langkah 2: Isi Template**
        - Isi kolom **Kode Unik** dengan Kode Unik FPTK yang valid
        - Isi kolom **Recruitment Update** dengan progress text
        - Isi kolom **Next Action** (kalau ada)

        **Langkah 3: Upload Kembali**
        - Upload file Excel yang sudah diisi di section **📤 Upload File Excel** di bawah
        - Sistem akan otomatis detect kolom dari template

        ### 📌 Tips
        - **Kode Unik** harus sama persis dengan yang ada di DB FPTK
        - Kolom lain (Posisi, PIC, BU, dll) otomatis di-link dari FPTK via Kode Unik
        - Kalau upload file sendiri (bukan template), bisa manual mapping kolom

        ### ⚠️ Penting
        - Kolom **Kode Unik** dan **Recruitment Update** wajib ada
        - Kode Unik yang gak ketemu di DB FPTK akan di-skip
        - Bisa upload multiple kandidat sekaligus
        """)

    # ============================================================
    # UPLOAD FILE
    # ============================================================
    st.markdown("---")
    st.markdown("#### 📤 Upload File Excel")

    col1, col2 = st.columns(2)
    with col1:
        week_num = st.number_input(
            "Week Number",
            min_value=1, max_value=53,
            value=current_week,
            key="up_upload_week"
        )
    with col2:
        year = st.number_input(
            "Year",
            min_value=2024, max_value=2030,
            value=current_year,
            key="up_upload_year"
        )

    uploaded = st.file_uploader(
        "Pilih file Excel",
        type=["xlsx", "xlsm"],
        key="up_upload_file"
    )

    if uploaded:
        try:
            xls = pd.ExcelFile(uploaded)
            sheet_names = xls.sheet_names

            selected_sheet = st.selectbox(
                "Pilih Sheet",
                sheet_names,
                index=sheet_names.index("COPAS yang ini") if "COPAS yang ini" in sheet_names else 0,
                key="up_upload_sheet"
            )

            df = pd.read_excel(uploaded, sheet_name=selected_sheet)

            st.success(f"✅ File terbaca: {len(df)} rows, {len(df.columns)} kolom")
            st.markdown("**Preview 5 rows:**")
            st.dataframe(df.head(5), use_container_width=True)

            # ============================================================
            # AUTO-DETECT KOLOM
            # ============================================================
            kolom_kode = None
            kolom_progress = None
            kolom_next_action = None

            for col in df.columns:
                cl = str(col).lower().strip()
                if "kode" in cl and "unik" in cl:
                    kolom_kode = col
                if "recruitment" in cl and "update" in cl:
                    kolom_progress = col
                elif "progress" in cl and not kolom_progress:
                    kolom_progress = col
                if "next" in cl and "action" in cl:
                    kolom_next_action = col

            # Cek apakah dari template (semua kolom ketemu)
            from_template = (kolom_kode is not None and kolom_progress is not None)

            st.markdown("---")

            if from_template:
                st.info("✅ Kolom terdeteksi otomatis! (kemungkinan dari template)")

                with st.expander("🔧 Ubah Mapping Kolom (opsional)", expanded=False):
                    col1, col2 = st.columns(2)
                    with col1:
                        kode_col = st.selectbox(
                            "Kolom **Kode Unik**",
                            df.columns.tolist(),
                            index=df.columns.tolist().index(kolom_kode) if kolom_kode in df.columns else 0,
                            key="up_map_kode"
                        )
                    with col2:
                        progress_col = st.selectbox(
                            "Kolom **Recruitment Update**",
                            df.columns.tolist(),
                            index=df.columns.tolist().index(kolom_progress) if kolom_progress in df.columns else 0,
                            key="up_map_progress"
                        )

                    next_action_col = st.selectbox(
                        "Kolom **Next Action** (opsional)",
                        ["(tidak ada)"] + df.columns.tolist(),
                        index=df.columns.tolist().index(kolom_next_action) + 1 if kolom_next_action in df.columns else 0,
                        key="up_map_next"
                    )
            else:
                st.warning("⚠️ Beberapa kolom tidak terdeteksi. Silakan mapping manual di bawah.")
                st.caption("💡 Kalau pakai template dari admin, kolom harus ada: **Kode Unik** & **Recruitment Update**")

                col1, col2 = st.columns(2)
                with col1:
                    kode_col = st.selectbox(
                        "Kolom **Kode Unik**",
                        df.columns.tolist(),
                        index=df.columns.tolist().index(kolom_kode) if kolom_kode in df.columns else 0,
                        key="up_map_kode_manual"
                    )
                with col2:
                    progress_col = st.selectbox(
                        "Kolom **Recruitment Update**",
                        df.columns.tolist(),
                        index=df.columns.tolist().index(kolom_progress) if kolom_progress in df.columns else 0,
                        key="up_map_progress_manual"
                    )

                next_action_col = st.selectbox(
                    "Kolom **Next Action** (opsional)",
                    ["(tidak ada)"] + df.columns.tolist(),
                    index=df.columns.tolist().index(kolom_next_action) + 1 if kolom_next_action in df.columns else 0,
                    key="up_map_next_manual"
                )

            st.markdown("---")
            st.markdown("#### 👀 Preview Data yang Akan Di-import")

            preview_rows = []
            for _, row in df.head(10).iterrows():
                kode = str(row.get(kode_col, '')).strip()
                progress = str(row.get(progress_col, '')).strip()
                if kode and kode != 'nan' and progress and progress != 'nan':
                    preview_rows.append({
                        "Kode Unik": kode,
                        "Progress (60 char)": progress[:60] + "..." if len(progress) > 60 else progress
                    })

            if preview_rows:
                st.dataframe(pd.DataFrame(preview_rows), use_container_width=True, hide_index=True)
            else:
                st.warning("⚠️ Gak ada row yang valid untuk di-import.")

            if st.button("🚀 Mulai Import", type="primary", use_container_width=True, key="up_btn_import"):
                if not preview_rows:
                    st.error("❌ Tidak ada data valid untuk di-import.")
                else:
                    progress_bar = st.progress(0)
                    status_text = st.empty()

                    created = 0
                    updated = 0
                    skipped_no_fptk = 0
                    skipped_empty = 0
                    errors = 0
                    error_details = []

                    total_rows = len(df)

                    for idx, row in df.iterrows():
                        progress_bar.progress((idx + 1) / total_rows)
                        status_text.info(f"Memproses row {idx + 1}/{total_rows}...")

                        kode_unik = str(row.get(kode_col, '')).strip()
                        progress_text = row.get(progress_col, '')

                        if not kode_unik or kode_unik == 'nan' or kode_unik == '':
                            skipped_empty += 1
                            continue

                        if pd.isna(progress_text) or str(progress_text).strip() == '':
                            skipped_empty += 1
                            continue

                        progress_text = str(progress_text).strip()

                        next_action_text = ""
                        if next_action_col and next_action_col != "(tidak ada)":
                            na = row.get(next_action_col, '')
                            if not pd.isna(na):
                                next_action_text = str(na).strip()

                        fptk = db.query(FPTK).filter(FPTK.kode_unik == kode_unik).first()

                        if not fptk:
                            skipped_no_fptk += 1
                            error_details.append(f"❌ Kode Unik '{kode_unik}' gak ada di DB FPTK")
                            continue

                        try:
                            progress, action = upsert_progress(
                                db, fptk, week_num, year,
                                progress_text, next_action_text
                            )
                            progress.created_by = user.id
                            progress.created_by_name = user.display_name or user.username

                            if action == "created":
                                created += 1
                            else:
                                updated += 1

                            if (created + updated) % 10 == 0:
                                db.commit()

                        except Exception as e:
                            errors += 1
                            error_details.append(f"❌ Row {idx+2}: {str(e)}")
                            db.rollback()

                    db.commit()
                    progress_bar.empty()
                    status_text.empty()

                    st.success(f"✅ Import selesai!")
                    st.markdown("### 📊 Summary")
                    col1, col2, col3, col4 = st.columns(4)
                    col1.metric("✅ Created", created)
                    col2.metric("🔄 Updated", updated)
                    col3.metric("⚠️ Skip (no FPTK)", skipped_no_fptk)
                    col4.metric("⚠️ Skip (empty)", skipped_empty)

                    if errors > 0:
                        st.error(f"❌ {errors} error saat import:")

                    if error_details:
                        with st.expander("🔍 Detail log"):
                            for d in error_details[:100]:
                                st.text(d)

                    st.cache_data.clear()

                    if created + updated > 0:
                        st.balloons()

        except Exception as e:
            st.error(f"❌ Error baca Excel: {str(e)}")
            import traceback
            with st.expander("🔍 Detail error"):
                st.code(traceback.format_exc())


# ============================================================
# TAB 3: EXPORT EXCEL
# ============================================================

def tab_export_excel(db, user, admin, current_week, current_year, filter_opts):
    st.markdown("### 📥 Export Excel Update Progress Recruitment")
    st.caption("Download data progress recruitment dalam format Excel yang siap dipakai.")

    st.markdown("#### 🔍 Filter Data yang Akan Di-export")

    col1, col2, col3 = st.columns(3)

    with col1:
        status_filter = st.selectbox(
            "Status Rekrutmen",
            ["OP", "Closed", "Cancel", "Semua"],
            index=0,
            key="export_status_filter"
        )

    with col2:
        pic_options = ["Semua"] + filter_opts.get("pic_options", [])
        pic_export = st.selectbox("PIC Recruiter", pic_options, key="export_pic_filter")

    with col3:
        bu_options = ["Semua"] + filter_opts.get("bu_options", [])
        bu_export = st.selectbox("Business Unit", bu_options, key="export_bu_filter")

    col4, col5 = st.columns(2)
    with col4:
        week_export = st.number_input(
            "Week Number", min_value=1, max_value=53,
            value=current_week, key="export_week_num"
        )
    with col5:
        year_export = st.number_input(
            "Year", min_value=2024, max_value=2030,
            value=current_year, key="export_year"
        )

    only_with_progress = st.checkbox(
        "Hanya tampilkan FPTK yang sudah ada progress di week ini",
        value=True,
        key="export_only_with_progress"
    )

    st.markdown("---")

    query = db.query(FPTK)

    if status_filter != "Semua":
        query = query.filter(FPTK.status == status_filter)
    if pic_export != "Semua":
        query = query.filter(FPTK.pic_recruiter == pic_export)
    if bu_export != "Semua":
        query = query.filter(FPTK.business_unit == bu_export)
    if not admin:
        query = query.filter(FPTK.pic_recruiter == user.pic_recruiter)

    fptk_list = query.order_by(FPTK.fptk_date_real.desc()).all()

    if not fptk_list:
        st.warning("⚠️ Tidak ada FPTK yang sesuai filter.")
        return

    progress_map = {}
    if only_with_progress:
        progress_records = db.query(RecruitmentProgress).filter(
            RecruitmentProgress.week_number == week_export,
            RecruitmentProgress.year == year_export
        ).all()
        progress_map = {p.fptk_id: p for p in progress_records}

    rows = []
    no = 1
    for fptk in fptk_list:
        progress = progress_map.get(fptk.id)

        if only_with_progress and not progress:
            continue

        progress_text = ""
        if progress:
            ptw = (progress.progress_this_week or "").strip()
            na = (progress.next_action or "").strip()

            if ptw.lower().startswith("progress weekly:"):
                ptw = ptw[len("progress weekly:"):].strip()
            if na.lower().startswith("next action:"):
                na = na[len("next action:"):].strip()

            if ptw:
                progress_text = f"Progress Weekly: {ptw}"
            if na:
                if progress_text:
                    progress_text += f"\nNext Action: {na}"
                else:
                    progress_text = f"Next Action: {na}"

        rows.append({
            'No': no,
            'Tanggal FPTK': fptk.fptk_date_real,
            'Posisi': fptk.posisi,
            'Level': fptk.level_fptk,
            'Business Unit': fptk.business_unit,
            'Filter Kategorisasi': fptk.filter_kategorisasi_fptk or fptk.category_fptk or "",
            'SLA Target Pemenuhan': fptk.deadline_sla,
            'PIC TA': fptk.pic_recruiter,
            'Vacancy': fptk.vacancy or 1,
            'Status Rekrutmen': fptk.status,
            'Recruitment Update': progress_text,
        })
        no += 1

    df_export = pd.DataFrame(rows)

    if df_export.empty:
        st.warning("⚠️ Tidak ada data untuk di-export.")
        return

    st.markdown(f"#### 👀 Preview ({len(df_export)} baris)")
    st.dataframe(df_export.head(20), use_container_width=True, hide_index=True)

    if len(df_export) > 20:
        st.caption(f"... dan {len(df_export) - 20} baris lainnya")

    st.markdown("---")

    week_label = get_week_label(week_export, year_export)

    col1, col2 = st.columns([1, 1])
    with col1:
        if st.button("📊 Generate Excel", type="primary", use_container_width=True, key="btn_generate_excel"):
            with st.spinner("Membuat file Excel..."):
                try:
                    excel_buffer = generate_progress_recruitment_excel(df_export)
                    st.session_state["excel_buffer"] = excel_buffer.getvalue()
                    st.session_state["excel_filename"] = (
                        f"Update_Progres_Recruitment_{week_label.replace(' ', '_').replace(',', '')}.xlsx"
                    )
                    st.success("✅ Excel berhasil di-generate!")
                except Exception as e:
                    st.error(f"❌ Error generate Excel: {str(e)}")
                    import traceback
                    with st.expander("🔍 Detail error"):
                        st.code(traceback.format_exc())

    with col2:
        if "excel_buffer" in st.session_state:
            st.download_button(
                label="⬇️ Download Excel",
                data=st.session_state["excel_buffer"],
                file_name=st.session_state["excel_filename"],
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
                use_container_width=True,
                key="btn_download_excel"
            )


# ============================================================
# MAIN FUNCTION
# ============================================================

def show_update_progres():
    st.title("📊 Update Progres Recruitment")
    st.markdown("Update progress rekrutmen mingguan per FPTK yang masih OP.")

    db = next(get_db())
    user = get_current_user(db)
    if not user:
        st.warning("Silakan login terlebih dahulu.")
        return

    admin = is_admin(db)
    current_week, current_year = get_current_week()
    week_label = get_week_label(current_week, current_year)
    week_start, week_end = get_week_range(current_week, current_year)

    filter_opts = _load_filter_options()

    col1, col2, col3 = st.columns(3)
    col1.metric("📅 Week Ini", week_label)
    if week_start and week_end:
        col2.metric("📆 Periode", f"{week_start.strftime('%d/%m')} - {week_end.strftime('%d/%m/%Y')}")
    col3.metric("👤 Login", user.display_name or user.username)

    st.markdown("---")

    with st.sidebar:
        st.markdown("### 🔍 Filter FPTK OP")

        st.text_input(
            "🔎 Cari (Kode Unik / Posisi)",
            key="search_progres",
            placeholder="Ketik keyword..."
        )

        pic_options = ["Semua"] + filter_opts.get("pic_options", [])
        default_pic_idx = 0
        if not admin and user.pic_recruiter in pic_options:
            default_pic_idx = pic_options.index(user.pic_recruiter)
        st.selectbox("PIC Recruiter", pic_options, index=default_pic_idx, key="pic_progres")

        st.selectbox("Business Unit", ["Semua"] + filter_opts.get("bu_options", []), key="bu_progres")
        st.selectbox("Direktorat", ["Semua"] + filter_opts.get("direktorat_options", []), key="dir_progres")
        st.selectbox("Divisi", ["Semua"] + filter_opts.get("divisi_options", []), key="divisi_progres")
        st.selectbox("Department", ["Semua"] + filter_opts.get("dept_options", []), key="dept_progres")
        st.selectbox("Filter Kategorisasi", ["Semua"] + filter_opts.get("filter_kategorisasi_options", []), key="kat_progres")
        st.selectbox("Level FPTK", ["Semua", "1A", "1B", "1C", "2A", "2B", "2C", "3A", "3B", "3C", "4A", "4B", "5A", "5B"], key="level_progres")

        st.markdown("---")
        st.checkbox("Hanya FPTK saya", value=not admin, key="show_mine_progres")

        if st.button("🔄 Refresh Filter Options", use_container_width=True, key="refresh_filter_progres"):
            _load_filter_options.clear()
            st.success("Filter refreshed!")
            time.sleep(0.3)
            st.rerun()

    tab1, tab2, tab3 = st.tabs([
        "✏️ Update Manual",
        "📤 Upload Excel Massal",
        "📥 Export Excel"
    ])

    with tab1:
        tab_update_manual(db, user, admin, current_week, current_year, filter_opts)

    with tab2:
        tab_upload_excel(db, user, admin, current_week, current_year)

    with tab3:
        tab_export_excel(db, user, admin, current_week, current_year, filter_opts)
